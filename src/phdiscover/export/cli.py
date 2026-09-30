"""
PhDiscover Engine - Export Module
Exports positions to Excel, JSON, CSV
"""

from __future__ import annotations

import asyncio
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import pandas as pd
import structlog

from phdiscover.config import get_settings
from phdiscover.models import Position
from phdiscover.models.discovery import ExportPosition

logger = structlog.get_logger(__name__)


class PositionExporter:
    """Exports positions to various formats"""

    def __init__(self):
        self.settings = get_settings()

    def to_excel(
        self,
        positions: List[Position],
        output_path: Path,
        min_score: float = 70.0,
    ) -> Path:
        """Export positions to Excel with formatting"""
        export_positions = [
            ExportPosition.from_position(p)
            for p in positions
            if p.final_rank_score >= min_score
        ]

        if not export_positions:
            logger.warning("no_positions_to_export")
            return output_path

        # Create DataFrame
        df = pd.DataFrame([p.model_dump() for p in export_positions])

        # Reorder columns for readability
        column_order = [
            "title", "opportunity_type", "university", "department", "country", "city",
            "pi_name", "pi_email", "application_deadline", "funding_amount",
            "funding_currency", "funding_details", "url", "source_url",
            "email_verified", "verification_level", "research_domain",
            "research_fit_score", "final_rank_score", "contact_priority",
            "discovered_at", "evidence_summary",
        ]
        df = df[column_order]

        # Write to Excel with formatting
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Positions")

            # Get workbook and worksheet for formatting
            workbook = writer.book
            worksheet = writer.sheets["Positions"]

            # Style header
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )

            for col_num, cell in enumerate(worksheet[1], 1):
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Set column widths
            column_widths = {
                "A": 50,  # title
                "B": 20,  # opportunity_type
                "C": 45,  # university
                "D": 35,  # department
                "E": 12,  # country
                "F": 20,  # city
                "G": 30,  # pi_name
                "H": 35,  # pi_email
                "I": 18,  # application_deadline
                "J": 18,  # funding_amount
                "K": 15,  # funding_currency
                "L": 30,  # funding_details
                "M": 55,  # url
                "N": 55,  # source_url
                "O": 15,  # email_verified
                "P": 18,  # verification_level
                "Q": 25,  # research_domain
                "R": 18,  # research_fit_score
                "S": 18,  # final_rank_score
                "T": 18,  # contact_priority
                "U": 22,  # discovered_at
                "V": 60,  # evidence_summary
            }

            for col_letter, width in column_widths.items():
                worksheet.column_dimensions[col_letter].width = width

            # Freeze top row
            worksheet.freeze_panes = "A2"

            # Add auto-filter
            worksheet.auto_filter.ref = worksheet.dimensions

            # Style data rows
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = thin_border

                    # Color code contact priority
                    if cell.column == 20:  # contact_priority column
                        if cell.value == "HIGH":
                            cell.fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
                        elif cell.value == "MEDIUM":
                            cell.fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")

        logger.info("excel_export_complete", path=str(output_path), count=len(export_positions))
        return output_path

    def to_json(
        self,
        positions: List[Position],
        output_path: Path,
        min_score: float = 70.0,
        include_evidence: bool = True,
    ) -> Path:
        """Export positions to JSON"""
        export_positions = [
            ExportPosition.from_position(p)
            for p in positions
            if p.final_rank_score >= min_score
        ]

        output_data = {
            "exported_at": datetime.utcnow().isoformat(),
            "total_positions": len(export_positions),
            "min_score_filter": min_score,
            "positions": [p.model_dump() for p in export_positions],
        }

        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(output_data, f, ensure_ascii=False, indent=2, default=str)

        logger.info("json_export_complete", path=str(output_path), count=len(export_positions))
        return output_path

    def to_csv(
        self,
        positions: List[Position],
        output_path: Path,
        min_score: float = 70.0,
    ) -> Path:
        """Export positions to CSV"""
        export_positions = [
            ExportPosition.from_position(p)
            for p in positions
            if p.final_rank_score >= min_score
        ]

        if not export_positions:
            logger.warning("no_positions_to_export")
            return output_path

        df = pd.DataFrame([p.model_dump() for p in export_positions])
        df.to_csv(output_path, index=False, encoding="utf-8")

        logger.info("csv_export_complete", path=str(output_path), count=len(export_positions))
        return output_path

    async def export_all_formats(
        self,
        positions: List[Position],
        base_path: Path,
        min_score: float = 70.0,
    ) -> Dict[str, Path]:
        """Export to all formats"""
        base_path.parent.mkdir(parents=True, exist_ok=True)

        results = {}
        results["excel"] = self.to_excel(positions, base_path.with_suffix(".xlsx"), min_score)
        results["json"] = self.to_json(positions, base_path.with_suffix(".json"), min_score)
        results["csv"] = self.to_csv(positions, base_path.with_suffix(".csv"), min_score)

        return results


async def main(
    format: str = "excel",
    output: Path = Path("exports/phdiscover_export"),
    min_score: float = 70.0,
    country: Optional[str] = None,
    days: int = 30,
):
    """CLI entry point"""
    # In production, load positions from database
    # For now, create sample data
    from phdiscover.models import Position, ContactPriority, OpportunityType
    from uuid import uuid4

    sample_positions = [
        Position(
            id=uuid4(),
            canonical_id="test_001",
            title="Fully Funded PhD in Biomechanics",
            opportunity_type=OpportunityType.FUNDED_PHD,
            university="University of Toronto",
            department="Kinesiology",
            country="CA",
            city="Toronto",
            pi_name="Dr. John Smith",
            pi_email="j.smith@utoronto.ca",
            application_deadline=datetime(2026, 1, 15).date(),
            funding_amount="35000",
            funding_currency="CAD",
            funding_details="NSERC Funded - $35,000/year",
            url="https://example.com/position/1",
            source_url="https://euraxess.ec.europa.eu/jobs/123",
            email_verified=True,
            verification_source="official_university",
            email_confidence=1.0,
            verification_level="STRONG",
            research_domain="biomechanics",
            research_fit_score=92.5,
            final_rank_score=94.2,
            contact_priority=ContactPriority.HIGH,
            discovered_at=datetime.utcnow(),
        ),
        Position(
            id=uuid4(),
            canonical_id="test_002",
            title="PhD Position in Motor Control",
            opportunity_type=OpportunityType.PHD_PROJECT,
            university="Delft University of Technology",
            department="Biomechanical Engineering",
            country="NL",
            city="Delft",
            pi_name="Prof. Maria van der Berg",
            pi_email="m.vandenberg@tudelft.nl",
            application_deadline=datetime(2026, 2, 1).date(),
            funding_amount="2800",
            funding_currency="EUR",
            funding_details="Fully funded - €2,800/month",
            url="https://example.com/position/2",
            source_url="https://www.findaphd.com/phds/456",
            email_verified=True,
            verification_source="official_university",
            email_confidence=1.0,
            verification_level="STRONG",
            research_domain="motor_control",
            research_fit_score=88.0,
            final_rank_score=89.5,
            contact_priority=ContactPriority.HIGH,
            discovered_at=datetime.utcnow(),
        ),
    ]

    exporter = PositionExporter()
    output.parent.mkdir(parents=True, exist_ok=True)

    if format == "excel":
        exporter.to_excel(sample_positions, output.with_suffix(".xlsx"), min_score)
    elif format == "json":
        exporter.to_json(sample_positions, output.with_suffix(".json"), min_score)
    elif format == "csv":
        exporter.to_csv(sample_positions, output.with_suffix(".csv"), min_score)
    else:
        await exporter.export_all_formats(sample_positions, output, min_score)

    logger.info("export_complete", format=format, output=str(output))


if __name__ == "__main__":
    asyncio.run(main())